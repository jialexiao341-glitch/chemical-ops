"""AI 附件解析器的无数据库单元测试。"""

from io import BytesIO

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pypdf import PdfWriter

from app.modules.files.parser import MAX_PARSE_BYTES, parse_file_content


def test_parse_text_and_keep_reference() -> None:
    result = parse_file_content("报价.csv", "text/csv", "供应商,单价\n甲公司,12.5".encode())

    assert result.status == "SUCCEEDED"
    assert "甲公司,12.5" in result.text
    assert result.references == [{"type": "text", "locator": "全文"}]


def test_parse_docx_xlsx_and_pptx_with_source_locators() -> None:
    document = Document()
    document.add_paragraph("安全库存说明")
    docx_buffer = BytesIO()
    document.save(docx_buffer)
    docx_result = parse_file_content("说明.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", docx_buffer.getvalue())

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "库存"
    sheet.append(["物料", "数量"])
    sheet.append(["乙醇", 12])
    xlsx_buffer = BytesIO()
    workbook.save(xlsx_buffer)
    xlsx_result = parse_file_content("库存.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", xlsx_buffer.getvalue())

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = slide.shapes.add_textbox(0, 0, 200, 100)
    textbox.text = "供应商评估"
    pptx_buffer = BytesIO()
    presentation.save(pptx_buffer)
    pptx_result = parse_file_content("评估.pptx", "application/vnd.openxmlformats-officedocument.presentationml.presentation", pptx_buffer.getvalue())

    assert "安全库存说明" in docx_result.text
    assert any(reference["type"] == "paragraph" for reference in docx_result.references)
    assert "乙醇" in xlsx_result.text
    assert any(reference["type"] == "sheet_row" for reference in xlsx_result.references)
    assert "供应商评估" in pptx_result.text
    assert pptx_result.references == [{"type": "slide", "slide": 1}]


def test_parse_pdf_and_reject_oversized_content() -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    pdf_buffer = BytesIO()
    writer.write(pdf_buffer)
    pdf_result = parse_file_content("空白.pdf", "application/pdf", pdf_buffer.getvalue())
    oversized_result = parse_file_content("超大.txt", "text/plain", b"x" * (MAX_PARSE_BYTES + 1))

    assert pdf_result.status == "EMPTY"
    assert oversized_result.status == "TOO_LARGE"

