"""受限的本地文档解析器，供 AI 附件上下文使用。"""

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from app.core.observability import metrics

MAX_PARSE_BYTES = 10 * 1024 * 1024
MAX_TEXT_CHARS = 40_000
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg"}
MAX_REFERENCE_COUNT = 100


@dataclass(frozen=True)
class ParsedFile:
    """解析结果，不包含原始二进制。"""

    status: str
    text: str
    parser: str
    references: list[dict[str, object]]
    error: str | None = None


def _trim(text: str) -> str:
    """规范化并限制送入模型的文本长度。"""

    normalized = "\n".join(line.rstrip() for line in text.replace("\r\n", "\n").replace("\r", "\n").splitlines()).strip()
    return normalized[:MAX_TEXT_CHARS]


def _result(parser: str, text: str, references: list[dict[str, object]]) -> ParsedFile:
    """构造成功、空内容或截断结果。"""

    cleaned = _trim(text)
    status = "SUCCEEDED" if cleaned else "EMPTY"
    metrics.increment("file_parse", labels={"parser": parser, "status": status})
    return ParsedFile(status=status, text=cleaned, parser=parser, references=references[:MAX_REFERENCE_COUNT])


def _parse_text(content: bytes, parser: str) -> ParsedFile:
    """解析 UTF-8 文本、CSV、Markdown 或 JSON。"""

    return _result(parser, content.decode("utf-8-sig", errors="replace"), [{"type": "text", "locator": "全文"}])


def _parse_pdf(content: bytes) -> ParsedFile:
    """解析 PDF 页面文本，并保留页码定位。"""

    from pypdf import PdfReader

    reader = PdfReader(BytesIO(content))
    chunks: list[str] = []
    references: list[dict[str, object]] = []
    for page_number, page in enumerate(reader.pages, start=1):
        page_text = page.extract_text() or ""
        if page_text.strip():
            chunks.append(f"[第 {page_number} 页]\n{page_text}")
            references.append({"type": "page", "page": page_number})
    result = _result("pypdf", "\n\n".join(chunks), references)
    if result.status == "EMPTY":
        return ParsedFile(status="OCR_REQUIRED", text="", parser="mineru-ocr-required", references=[], error="PDF没有可提取的文本，需要启用 MinerU 进行 OCR")
    return result


def _parse_docx(content: bytes) -> ParsedFile:
    """解析 DOCX 段落和表格。"""

    from docx import Document

    document = Document(BytesIO(content))
    chunks: list[str] = []
    references: list[dict[str, object]] = []
    for paragraph_number, paragraph in enumerate(document.paragraphs, start=1):
        if paragraph.text.strip():
            chunks.append(f"[段落 {paragraph_number}]\n{paragraph.text}")
            references.append({"type": "paragraph", "paragraph": paragraph_number})
    for table_number, table in enumerate(document.tables, start=1):
        rows = ["\t".join((cell.text or "").strip() for cell in row.cells) for row in table.rows]
        table_text = "\n".join(row for row in rows if row.strip())
        if table_text:
            chunks.append(f"[表格 {table_number}]\n{table_text}")
            references.append({"type": "table", "table": table_number})
    return _result("python-docx", "\n\n".join(chunks), references)


def _parse_xlsx(content: bytes) -> ParsedFile:
    """解析 XLSX 工作表中的非空行。"""

    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    chunks: list[str] = []
    references: list[dict[str, object]] = []
    try:
        for sheet in workbook.worksheets:
            rows: list[str] = []
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                values = [str(value).strip() for value in row if value is not None and str(value).strip()]
                if values:
                    rows.append(f"第 {row_number} 行：" + "\t".join(values))
                    references.append({"type": "sheet_row", "sheet": sheet.title, "row": row_number})
            if rows:
                chunks.append(f"[工作表：{sheet.title}]\n" + "\n".join(rows))
    finally:
        workbook.close()
    return _result("openpyxl", "\n\n".join(chunks), references)


def _parse_pptx(content: bytes) -> ParsedFile:
    """解析 PPTX 幻灯片中的文本。"""

    from pptx import Presentation

    presentation = Presentation(BytesIO(content))
    chunks: list[str] = []
    references: list[dict[str, object]] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        texts = [shape.text.strip() for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()]
        if texts:
            chunks.append(f"[第 {slide_number} 页]\n" + "\n".join(texts))
            references.append({"type": "slide", "slide": slide_number})
    return _result("python-pptx", "\n\n".join(chunks), references)


def parse_file_content(file_name: str, mime_type: str, content: bytes) -> ParsedFile:
    """按扩展名解析文件；不支持或超限时返回明确状态，不猜测正文。"""

    if len(content) > MAX_PARSE_BYTES:
        return ParsedFile(status="TOO_LARGE", text="", parser="none", references=[], error="文件超过 AI 单次解析上限")
    extension = Path(file_name).suffix.lower()
    if extension in {".txt", ".csv", ".md", ".json"} or mime_type.startswith("text/"):
        return _parse_text(content, extension.lstrip(".") or "text")
    if extension in IMAGE_EXTENSIONS or mime_type.startswith("image/"):
        return ParsedFile(status="OCR_REQUIRED", text="", parser="mineru-ocr-required", references=[], error="图片需要启用 MinerU 进行 OCR")
    if extension == ".pdf" or mime_type == "application/pdf":
        return _parse_pdf(content)
    if extension == ".docx":
        return _parse_docx(content)
    if extension == ".xlsx":
        return _parse_xlsx(content)
    if extension == ".pptx":
        return _parse_pptx(content)
    return ParsedFile(status="UNSUPPORTED", text="", parser="none", references=[], error=f"暂不支持解析 {extension or mime_type or '未知'} 文件")
